"""Downloadable Word and PDF versions of an evidence-grounded Markdown report.

The report content is produced by :mod:`teren_oi.report`. This module only handles
presentation; it does not reinterpret findings or remove source citations.
"""

from __future__ import annotations

import html
import os
from string import punctuation
from io import BytesIO
from pathlib import Path
from threading import Lock
from typing import Literal

_MAX_REPORT_LENGTH = 1_000_000
_FONT_LOCK = Lock()
_PDF_FONT_NAMES: tuple[str, str] | None = None


def _lines(report_markdown: str) -> list[tuple[str, str]]:
    """Reduce the report's limited Markdown vocabulary to document blocks."""
    blocks: list[tuple[str, str]] = []
    for raw_line in report_markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("### "):
            blocks.append(("heading3", line[4:]))
        elif line.startswith("## "):
            blocks.append(("heading2", line[3:]))
        elif line.startswith("# "):
            blocks.append(("heading1", line[2:]))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:]))
        elif line.startswith("> "):
            blocks.append(("quote", line[2:]))
        else:
            blocks.append(("body", line))
    return blocks


def _inline_runs(markdown_text: str) -> list[tuple[str, bool]]:
    """Parse authored bold and literal escapes without reinterpreting source text.

    Escapes are decoded during scanning, so an escaped ``*`` never becomes a
    bold delimiter on a later pass. An unmatched authored delimiter stays literal.
    """
    tokens: list[tuple[str, bool]] = []
    text: list[str] = []
    index = 0
    while index < len(markdown_text):
        character = markdown_text[index]
        if character == "\\" and index + 1 < len(markdown_text) and markdown_text[index + 1] in punctuation:
            text.append(markdown_text[index + 1])
            index += 2
        elif markdown_text.startswith("**", index):
            if text:
                tokens.append(("".join(text), False))
                text = []
            tokens.append(("**", True))
            index += 2
        else:
            text.append(character)
            index += 1
    if text:
        tokens.append(("".join(text), False))
    markers = [index for index, (_, marker) in enumerate(tokens) if marker]
    paired = set(markers[:len(markers) - len(markers) % 2])
    runs: list[tuple[str, bool]] = []
    bold = False
    for index, (value, marker) in enumerate(tokens):
        if marker and index in paired:
            bold = not bold
        else:
            runs.append((value, bold))
    return runs


def _add_docx_runs(paragraph: object, markdown_text: str) -> None:
    for value, is_bold in _inline_runs(markdown_text):
        run = paragraph.add_run(value)
        run.bold = is_bold


def _export_docx(blocks: list[tuple[str, str]]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    document.core_properties.title = "Аналитическое заключение Teren Oi"
    document.core_properties.author = "Teren Oi"
    document.core_properties.subject = "Сопоставление редакций документов и проверяемые источники"
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.footer_distance = Cm(1)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    footer_run = footer.add_run("Teren Oi · Страница ")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(100, 116, 139)
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for heading_name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        style = document.styles[heading_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(31, 51, 84)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    for kind, value in blocks:
        if kind.startswith("heading"):
            level = int(kind[-1])
            paragraph = document.add_heading(level=level)
        elif kind == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
        else:
            paragraph = document.add_paragraph()
            if kind == "quote":
                paragraph.paragraph_format.left_indent = Cm(0.6)
                paragraph.paragraph_format.right_indent = Cm(0.3)
                paragraph.paragraph_format.space_after = Pt(9)
        _add_docx_runs(paragraph, value)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _font_candidates() -> list[tuple[Path, Path | None]]:
    package_fonts = Path(__file__).resolve().parent / "fonts"
    windows_fonts = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    return [
        (package_fonts / "DejaVuSans.ttf", package_fonts / "DejaVuSans-Bold.ttf"),
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
        (
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
        ),
        (windows_fonts / "arial.ttf", windows_fonts / "arialbd.ttf"),
        (windows_fonts / "calibri.ttf", windows_fonts / "calibrib.ttf"),
        (Path("/Library/Fonts/Arial.ttf"), Path("/Library/Fonts/Arial Bold.ttf")),
        (
            Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
            Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
        ),
    ]


def _pdf_fonts() -> tuple[str, str]:
    """Register a TrueType family that can render Cyrillic text."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    global _PDF_FONT_NAMES
    with _FONT_LOCK:
        if _PDF_FONT_NAMES is not None:
            return _PDF_FONT_NAMES
        for regular_path, bold_path in _font_candidates():
            if not regular_path.is_file():
                continue
            regular_name = "TerenOi-Regular"
            bold_name = "TerenOi-Bold"
            pdfmetrics.registerFont(TTFont(regular_name, str(regular_path)))
            if bold_path is not None and bold_path.is_file():
                pdfmetrics.registerFont(TTFont(bold_name, str(bold_path)))
            else:
                bold_name = regular_name
            pdfmetrics.registerFontFamily(
                regular_name,
                normal=regular_name,
                bold=bold_name,
                italic=regular_name,
                boldItalic=bold_name,
            )
            _PDF_FONT_NAMES = (regular_name, bold_name)
            return _PDF_FONT_NAMES
    raise RuntimeError(
        "PDF export requires a Cyrillic TrueType font (Arial, Calibri, DejaVu Sans, "
        "or Liberation Sans). Install one system-wide or bundle DejaVuSans.ttf "
        "under src/teren_oi/fonts/."
    )


def _pdf_markup(markdown_text: str) -> str:
    parts = []
    for value, is_bold in _inline_runs(markdown_text):
        escaped = html.escape(value, quote=False)
        parts.append(f"<b>{escaped}</b>" if is_bold else escaped)
    return "".join(parts)


def _export_pdf(blocks: list[tuple[str, str]]) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate
    except ImportError as exc:
        raise RuntimeError("PDF export requires the reportlab package") from exc

    regular_font, bold_font = _pdf_fonts()
    styles = {
        "body": ParagraphStyle(
            "TerenOiBody",
            fontName=regular_font,
            fontSize=9.5,
            leading=14,
            alignment=TA_LEFT,
            spaceAfter=8,
            splitLongWords=True,
        ),
        "heading1": ParagraphStyle(
            "TerenOiHeading1",
            fontName=bold_font,
            fontSize=17,
            leading=21,
            textColor=colors.HexColor("#1F3354"),
            spaceBefore=5,
            spaceAfter=14,
            keepWithNext=True,
        ),
        "heading2": ParagraphStyle(
            "TerenOiHeading2",
            fontName=bold_font,
            fontSize=12.5,
            leading=17,
            textColor=colors.HexColor("#1F3354"),
            spaceBefore=14,
            spaceAfter=8,
            keepWithNext=True,
        ),
        "heading3": ParagraphStyle(
            "TerenOiHeading3",
            fontName=bold_font,
            fontSize=10.5,
            leading=15,
            textColor=colors.HexColor("#1F3354"),
            spaceBefore=10,
            spaceAfter=6,
            keepWithNext=True,
        ),
        "bullet": ParagraphStyle(
            "TerenOiBullet",
            fontName=regular_font,
            fontSize=9.5,
            leading=14,
            leftIndent=12,
            firstLineIndent=-9,
            spaceAfter=6,
            splitLongWords=True,
        ),
        "quote": ParagraphStyle(
            "TerenOiQuote",
            fontName=regular_font,
            fontSize=9,
            leading=13,
            leftIndent=15,
            rightIndent=8,
            textColor=colors.HexColor("#475569"),
            spaceAfter=9,
            splitLongWords=True,
        ),
    }
    story = []
    for kind, value in blocks:
        content = _pdf_markup(value)
        if kind == "bullet":
            content = f"• {content}"
        story.append(Paragraph(content, styles[kind]))

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=19 * mm,
        rightMargin=19 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title="Аналитическое заключение Teren Oi",
        author="Teren Oi",
    )

    def page_footer(canvas: object, template: object) -> None:
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#E2E8F0"))
        canvas.setLineWidth(0.5)
        canvas.line(19 * mm, 15 * mm, A4[0] - 19 * mm, 15 * mm)
        canvas.setFont(regular_font, 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(19 * mm, 10.5 * mm, "Teren Oi")
        canvas.drawRightString(A4[0] - 19 * mm, 10.5 * mm, f"Страница {template.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return output.getvalue()


def export_report(
    report_markdown: str, format: Literal["pdf", "docx"]
) -> tuple[bytes, str, str]:
    """Render a report and return (content, media type, download filename)."""
    if not isinstance(report_markdown, str) or not report_markdown.strip():
        raise ValueError("Report must contain non-empty Markdown text")
    if len(report_markdown) > _MAX_REPORT_LENGTH:
        raise ValueError(f"Report exceeds {_MAX_REPORT_LENGTH} characters")
    blocks = _lines(report_markdown)
    if format == "docx":
        return (
            _export_docx(blocks),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "teren_oi_report.docx",
        )
    if format == "pdf":
        return _export_pdf(blocks), "application/pdf", "teren_oi_report.pdf"
    raise ValueError("Unsupported report format; expected 'pdf' or 'docx'")
