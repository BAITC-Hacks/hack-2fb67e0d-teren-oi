from __future__ import annotations

import json
import unittest
from io import BytesIO
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi.testclient import TestClient
from openpyxl import Workbook

from teren_oi import web
from teren_oi.analyzer import _MAX_EVIDENCE_CHARS, _payload_evidence, evidence_coverage
from teren_oi.diff import compare_documents
from teren_oi.docx_reader import read_docx
from teren_oi.document_safety import DocumentReadError
from teren_oi.models import Clause, SourceDocument
from teren_oi.parsers import TextBlock, parse_blocks
from teren_oi.readers import read_txt, read_xlsx


class AuditRegressionTests(unittest.TestCase):
    def test_direct_lettered_ids_and_following_subpoint(self):
        result = parse_blocks([TextBlock(
            "3.4.а Проверяет сеть\nб) Готовит отчёт\n3.4.v Контролирует сеть", "line 1"
        )], "input.txt")
        self.assertEqual([c.clause_id for c in result.clauses], ["3.4.а", "3.4.б", "3.4.v"])
        self.assertEqual(result.clauses[0].text, "Проверяет сеть")

    def test_numeric_prose_is_not_a_clause(self):
        for text in ("2026год работы", "100% заявок проверяется", "3.14% ошибок", "9" * 5000 + " текст"):
            with self.subTest(text=text[:25]):
                result = parse_blocks([TextBlock(text, "line 1")], "input.txt")
                self.assertFalse(result.clauses)
                self.assertEqual(result.unnumbered_blocks, (text,))

    def test_standalone_compound_id_keeps_correct_parent(self):
        result = parse_blocks([TextBlock("3.4.а", "1"), TextBlock("Проверяет сеть", "2"),
                               TextBlock("б) Готовит отчёт", "3")], "input.txt")
        self.assertEqual([c.clause_id for c in result.clauses], ["3.4.а", "3.4.б"])

    def test_numbered_punctuation_does_not_require_a_space(self):
        result = parse_blocks([TextBlock("1)Проверяет сеть\n2.Готовит отчёт", "1")], "input.txt")
        self.assertEqual([c.clause_id for c in result.clauses], ["1", "2"])

    def test_serialized_ai_payload_stays_within_budget(self):
        for unchanged, body in ((False, "x"), (False, '"\\\n' * 1000), (True, "x")):
            with self.subTest(unchanged=unchanged, escaping=len(body) > 1):
                clauses = tuple(Clause(str(i), body, "test", str(i)) for i in range(1000))
                comparison = compare_documents(SourceDocument("before", clauses if unchanged else ()),
                                               SourceDocument("after", clauses))
                parts, _ = _payload_evidence(comparison)
                serialized = json.dumps({"context_complete": {"до": False, "после": False},
                                         "clauses": parts}, ensure_ascii=False)
                self.assertLessEqual(len(serialized), _MAX_EVIDENCE_CHARS)
                self.assertEqual(evidence_coverage(comparison, sent=True)["included_clauses"], len(parts))
                self.assertGreater(len(parts), 0)

    def test_office_archive_expansion_is_bounded(self):
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as archive:
            archive.writestr("huge.xml", "x" * 1024)
        for reader, name in ((read_docx, "big.docx"), (read_xlsx, "big.xlsx")):
            with self.subTest(name=name), patch("teren_oi.document_safety.MAX_ARCHIVE_BYTES", 512):
                with self.assertRaisesRegex(DocumentReadError, "распакованный"):
                    reader(output.getvalue(), name)

    def test_xlsx_huge_sparse_dimensions_are_rejected(self):
        workbook = Workbook()
        workbook.active["A1"] = "1.1 Проверяет сеть"
        workbook.active["XFD1048576"] = "tail"
        output = BytesIO()
        workbook.save(output)
        with self.assertRaisesRegex(DocumentReadError, "500 000 ячеек"):
            read_xlsx(output.getvalue(), "sparse.xlsx")

    def test_invalid_controls_are_rejected_in_file_text_and_export(self):
        with self.assertRaises(DocumentReadError):
            read_txt(b"1.1 Bad\x00text", "bad.txt")
        client = TestClient(web.app)
        response = client.post("/api/analyze", data={"before_text": "1.1 Bad\x00text", "after_text": "1.1 Good"})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["error"]["code"], "INVALID_TEXT")
        response = client.post("/api/export", json={"report_markdown": "Bad\x00text", "format": "docx"})
        self.assertEqual(response.status_code, 400)

    def test_local_signal_is_separate_from_ai_loss_in_report(self):
        payload = TestClient(web.app).post("/api/analyze", data={"demo": "true"}).json()
        report = payload["report_markdown"]
        loss_section = report.split("## Potential Lost Functions", 1)[1].split("## Local Signals", 1)[0]
        self.assertNotIn(payload["findings"][0]["title"], loss_section)
        self.assertIn(payload["findings"][0]["title"], report.split("## Local Signals", 1)[1])

    def test_two_valid_large_inputs_allow_report_over_one_million_chars(self):
        client = TestClient(web.app)
        response = client.post("/api/analyze", data={
            "before_text": "1.1 " + "A" * 499_990,
            "after_text": "1.1 " + "B" * 499_990,
        })
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertGreater(len(payload["report_markdown"]), 1_000_000)
        exported = client.post("/api/export", json={"analysis_id": payload["analysis_id"], "format": "docx"})
        self.assertEqual(exported.status_code, 200)
        self.assertTrue(exported.content.startswith(b"PK"))


if __name__ == "__main__":
    unittest.main()
