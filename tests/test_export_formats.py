from __future__ import annotations

import unittest
from io import BytesIO

import pymupdf
from docx import Document

from teren_oi.export_formats import export_report
from teren_oi.diff import compare_documents
from teren_oi.models import Citation, Clause, Finding, SourceDocument
from teren_oi.report import report_as_markdown


QUOTE = "Проверяет качество обслуживания и передаёт заключение руководителю."
SOURCE = "Положение_департамента_клиентской_аналитики_" * 4 + "2026.docx"


def sample_report() -> str:
    sections = [
        "# Аналитическое заключение Tereñ oi", "",
        f"- До: **{SOURCE}**", "- После: **Новая редакция.docx**", "",
        "## Итоговая сводка", "",
        "Изменения требуют ручной проверки. Данные относятся к синтетическому примеру.",
    ]
    for index in range(1, 35):
        sections.extend([
            f"## Замечание {index}", "",
            "Функция могла быть передана другому подразделению. " * 8,
            f"> До — {SOURCE}, абзац {index}, пункт 4.{index}: «{QUOTE}»",
            "Точные значения: 5 < 10, A&B. Конец фрагмента.",
        ])
    sections.extend(["## Ограничения", "Проверяйте выводы по исходным документам."])
    return "\n".join(sections)


class ExportFormatTests(unittest.TestCase):
    def test_literal_evidence_survives_both_exports_and_cannot_create_headings(self) -> None:
        quote = r"Проверяет **важные** документы C:\docs\* и \\server\folder."
        old = SourceDocument("old.txt", (Clause("1", quote, r"C:\docs\old.txt", "line 1"),))
        new = SourceDocument("new.txt", ())
        finding = Finding(
            kind="другое изменение", title="Вывод\n## Ложный заголовок",
            explanation="Первая строка.\n## Только текст модели\nВторая строка.",
            confidence="низкая",
            citations=[Citation(document_label="до", clause_id="1", quote=quote)],
        )
        markdown = report_as_markdown(compare_documents(old, new), [finding], "old.txt", "new.txt")
        for format in ("pdf", "docx"):
            with self.subTest(format=format):
                data, _, _ = export_report(markdown, format)
                if format == "pdf":
                    with pymupdf.open(stream=data, filetype="pdf") as pdf:
                        text = "\n".join(page.get_text() for page in pdf)
                else:
                    word = Document(BytesIO(data))
                    text = "\n".join(p.text for p in word.paragraphs)
                    injected = [p for p in word.paragraphs if p.text == "## Только текст модели"]
                    self.assertTrue(injected)
                    self.assertTrue(all(p.style.name == "Normal" for p in injected))
                self.assertIn("".join(quote.split()), "".join(text.split()))
                self.assertIn(r"C:\docs\old.txt", text)
                self.assertIn("Вторая строка.", text)

    def test_authored_bold_and_escaped_literals_can_share_a_run(self) -> None:
        markdown = r"**Метка \*\*важно\*\***: C:\\docs\\ и \\*"
        data, _, _ = export_report(markdown, "docx")
        paragraph = Document(BytesIO(data)).paragraphs[0]
        self.assertEqual(paragraph.text, r"Метка **важно**: C:\docs\ и \*")
        self.assertEqual("".join(run.text for run in paragraph.runs if run.bold), "Метка **важно**")

    def test_pdf_retains_cyrillic_sources_quotes_and_page_numbers(self) -> None:
        data, media_type, filename = export_report(sample_report(), "pdf")
        self.assertEqual((media_type, filename), ("application/pdf", "teren_oi_report.pdf"))
        with pymupdf.open(stream=data, filetype="pdf") as document:
            self.assertGreater(document.page_count, 3)
            self.assertEqual(document.metadata["author"], "Tereñ oi")
            text = "\n".join(page.get_text() for page in document)
            compact = "".join(text.split())
            self.assertIn("".join(QUOTE.split()), compact)
            self.assertIn(SOURCE, compact)
            self.assertIn("5<10,A&B.", compact)
            self.assertIn("Проверяйтевыводыпоисходнымдокументам.", compact)
            for number, page in enumerate(document, start=1):
                self.assertIn(f"Страница {number}", page.get_text())
                for block in page.get_text("blocks"):
                    self.assertGreaterEqual(block[0], 0)
                    self.assertGreaterEqual(block[1], 0)
                    self.assertLessEqual(block[2], page.rect.width + 1)
                    self.assertLessEqual(block[3], page.rect.height + 1)

    def test_word_retains_cyrillic_source_quotes_and_uses_dynamic_pagination(self) -> None:
        data, media_type, filename = export_report(sample_report(), "docx")
        self.assertEqual(filename, "teren_oi_report.docx")
        self.assertIn("wordprocessingml", media_type)
        document = Document(BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn(QUOTE, text)
        self.assertIn(SOURCE, text)
        self.assertIn("5 < 10, A&B.", text)
        self.assertIn("Замечание 34", text)
        self.assertEqual(document.core_properties.author, "Tereñ oi")
        self.assertIn("PAGE", document.sections[0].footer._element.xml)
        self.assertAlmostEqual(document.sections[0].page_width.mm, 210, delta=0.1)
        self.assertAlmostEqual(document.sections[0].page_height.mm, 297, delta=0.1)

    def test_long_single_quote_can_split_across_pdf_pages_without_losing_tail(self) -> None:
        report = "# Источник\n\n> " + ("Ответственное подразделение проверяет данные. " * 450)
        report += "ПоследняяФразаИсточника"
        data, _, _ = export_report(report, "pdf")
        with pymupdf.open(stream=data, filetype="pdf") as document:
            self.assertGreater(document.page_count, 1)
            self.assertIn("ПоследняяФразаИсточника", document[-1].get_text())


if __name__ == "__main__":
    unittest.main()
