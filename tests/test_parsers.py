from __future__ import annotations

import unittest
from io import BytesIO

import pymupdf
from docx import Document
from openpyxl import Workbook

from teren_oi.docx_reader import DocumentReadError, read_docx
from teren_oi.parsers import TextBlock, parse_blocks
from teren_oi.readers import read_document, read_pdf, read_txt, read_xlsx


class BlockParserTests(unittest.TestCase):
    def test_nested_and_lettered_clauses_keep_parent_and_location(self) -> None:
        document = parse_blocks(
            [
                TextBlock(
                    "5.3.2 Функции подразделения:\nа) анализ данных\nб. контроль качества",
                    "page 4 / block 2",
                )
            ],
            "regulation.pdf",
        )

        self.assertEqual(
            [clause.clause_id for clause in document.clauses],
            ["5.3.2", "5.3.2.а", "5.3.2.б"],
        )
        self.assertEqual(document.clauses[1].text, "анализ данных")
        self.assertEqual(document.clauses[1].source, "regulation.pdf")
        self.assertIn("page 4", document.clauses[1].location)
        self.assertIn("subpoint а", document.clauses[1].location)

    def test_identifier_in_separate_block_uses_next_block_as_body(self) -> None:
        document = parse_blocks(
            [TextBlock("7.2.1", "block 1"), TextBlock("Проводит аудит", "block 2")],
            "source.docx",
        )

        self.assertEqual(len(document.clauses), 1)
        self.assertEqual(document.clauses[0].clause_id, "7.2.1")
        self.assertEqual(document.clauses[0].text, "Проводит аудит")
        self.assertIn("continuation block 2", document.clauses[0].location)

    def test_lettered_identifier_in_separate_block_uses_numeric_parent(self) -> None:
        document = parse_blocks(
            [
                TextBlock("7.2 Общие функции", "block 1"),
                TextBlock("б)", "block 2"),
                TextBlock("Контролирует качество", "block 3"),
            ],
            "source.docx",
        )

        self.assertEqual(document.clauses[1].clause_id, "7.2.б")
        self.assertEqual(document.clauses[1].text, "Контролирует качество")
        self.assertIn("subpoint б", document.clauses[1].location)


class ReaderTests(unittest.TestCase):
    def test_docx_preserves_paragraph_table_order_and_lettered_subpoints(self) -> None:
        source = Document()
        source.add_paragraph("Заголовок")
        source.add_paragraph("5.3.2 Основная функция")
        source.add_paragraph("а) Анализ данных")
        table = source.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "6.1"
        table.cell(0, 1).text = "Табличная функция"
        stream = BytesIO()
        source.save(stream)

        document = read_docx(stream.getvalue(), "audit.docx")

        self.assertEqual(
            [clause.clause_id for clause in document.clauses],
            ["5.3.2", "5.3.2.а", "6.1"],
        )
        self.assertEqual(document.unnumbered_blocks, ("Заголовок",))
        self.assertIn("paragraph", document.clauses[0].location)
        self.assertIn("table 1 / row 1", document.clauses[2].location)

    def test_pdf_reads_text_layer_and_reports_page(self) -> None:
        source = pymupdf.open()
        page = source.new_page()
        page.insert_text((72, 72), "1.1 Audit function")
        data = source.tobytes()
        source.close()

        document = read_pdf(data, "audit.pdf")

        self.assertEqual(document.clauses[0].clause_id, "1.1")
        self.assertEqual(document.clauses[0].text, "Audit function")
        self.assertIn("page 1", document.clauses[0].location)

    def test_pdf_without_text_explains_ocr_limit(self) -> None:
        source = pymupdf.open()
        source.new_page()
        data = source.tobytes()
        source.close()

        with self.assertRaisesRegex(DocumentReadError, r"OCR.*недоступно.*Загрузите PDF"):
            read_pdf(data, "scan.pdf")

    def test_xlsx_combines_clause_id_and_body_cells(self) -> None:
        source = Workbook()
        sheet = source.active
        sheet.title = "Functions"
        sheet["A1"] = "5.3.2"
        sheet["B1"] = "Анализ данных"
        sheet["A2"] = "а)"
        sheet["B2"] = "Контроль качества"
        stream = BytesIO()
        source.save(stream)

        document = read_xlsx(stream.getvalue(), "functions.xlsx")

        self.assertEqual(
            [clause.clause_id for clause in document.clauses],
            ["5.3.2", "5.3.2.а"],
        )
        self.assertEqual(document.clauses[0].text, "Анализ данных")
        self.assertIn("Functions!A1", document.clauses[0].location)

    def test_txt_and_dispatcher_are_case_insensitive(self) -> None:
        document = read_document("2.1 Проверка".encode(), "RULES.TXT")
        direct = read_txt("2.1 Проверка".encode(), "RULES.TXT")

        self.assertEqual(document, direct)
        self.assertEqual(document.clauses[0].location, "line 1 / section 2.1")

    def test_dispatcher_rejects_unknown_extension(self) -> None:
        with self.assertRaisesRegex(DocumentReadError, "не поддерживается"):
            read_document(b"data", "rules.csv")


if __name__ == "__main__":
    unittest.main()
